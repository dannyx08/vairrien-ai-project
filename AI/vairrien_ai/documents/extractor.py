from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Any

from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation


class DocumentExtractor:
    """
    Vairrien document extraction layer.

    Converts supported files into a normalized structure
    that the Vairrien AI system can process.

    Supported formats:
        TXT
        PDF
        DOCX
        CSV
        XLSX
        PPTX
    """

    SUPPORTED_EXTENSIONS = {
        ".txt",
        ".pdf",
        ".docx",
        ".csv",
        ".xlsx",
        ".pptx",
    }

    async def extract(
        self,
        file_path: str | Path,
    ) -> dict[str, Any]:
        """
        Extract content from a supported document.
        """

        path = Path(file_path)

        if not path.exists():
            return {
                "status": "error",
                "error": "File does not exist.",
            }

        if not path.is_file():
            return {
                "status": "error",
                "error": "Path is not a file.",
            }

        extension = path.suffix.lower()

        if extension not in self.SUPPORTED_EXTENSIONS:
            return {
                "status": "error",
                "error": f"Unsupported file type: {extension}",
                "supported_types": sorted(
                    self.SUPPORTED_EXTENSIONS
                ),
            }

        try:
            if extension == ".txt":
                result = self._extract_txt(path)

            elif extension == ".pdf":
                result = self._extract_pdf(path)

            elif extension == ".docx":
                result = self._extract_docx(path)

            elif extension == ".csv":
                result = self._extract_csv(path)

            elif extension == ".xlsx":
                result = self._extract_xlsx(path)

            elif extension == ".pptx":
                result = self._extract_pptx(path)

            else:
                return {
                    "status": "error",
                    "error": "Unsupported file type.",
                }

            return {
                "status": "success",
                "filename": path.name,
                "file_type": extension.lstrip("."),
                **result,
            }

        except Exception as exc:
            return {
                "status": "error",
                "filename": path.name,
                "file_type": extension.lstrip("."),
                "error": str(exc),
            }

    # =========================================================
    # TXT
    # =========================================================

    def _extract_txt(
        self,
        path: Path,
    ) -> dict[str, Any]:

        text = path.read_text(
            encoding="utf-8",
            errors="replace",
        )

        return {
            "content": text,
            "metadata": {
                "characters": len(text),
                "lines": len(text.splitlines()),
            },
        }

    # =========================================================
    # PDF
    # =========================================================

    def _extract_pdf(
        self,
        path: Path,
    ) -> dict[str, Any]:

        reader = PdfReader(str(path))

        pages = []

        for page_number, page in enumerate(
            reader.pages,
            start=1,
        ):
            text = page.extract_text() or ""

            pages.append(
                {
                    "page": page_number,
                    "content": text,
                }
            )

        full_text = "\n\n".join(
            page["content"]
            for page in pages
        )

        return {
            "content": full_text,
            "pages": pages,
            "metadata": {
                "page_count": len(reader.pages),
                "characters": len(full_text),
            },
        }

    # =========================================================
    # DOCX
    # =========================================================

    def _extract_docx(
        self,
        path: Path,
    ) -> dict[str, Any]:

        document = Document(str(path))

        paragraphs = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()

            if text:
                paragraphs.append(text)

        tables = []

        for table in document.tables:
            rows = []

            for row in table.rows:
                rows.append(
                    [
                        cell.text.strip()
                        for cell in row.cells
                    ]
                )

            tables.append(rows)

        content_parts = []

        if paragraphs:
            content_parts.append(
                "\n".join(paragraphs)
            )

        for table_number, table in enumerate(
            tables,
            start=1,
        ):
            table_text = "\n".join(
                " | ".join(row)
                for row in table
            )

            content_parts.append(
                f"Table {table_number}:\n"
                f"{table_text}"
            )

        content = "\n\n".join(
            content_parts
        )

        return {
            "content": content,
            "paragraphs": paragraphs,
            "tables": tables,
            "metadata": {
                "paragraph_count": len(paragraphs),
                "table_count": len(tables),
                "characters": len(content),
            },
        }

    # =========================================================
    # CSV
    # =========================================================

    def _extract_csv(
        self,
        path: Path,
    ) -> dict[str, Any]:

        text = path.read_text(
            encoding="utf-8-sig",
            errors="replace",
        )

        sample = text[:4096]

        try:
            dialect = csv.Sniffer().sniff(
                sample
            )
        except csv.Error:
            dialect = csv.excel

        reader = csv.reader(
            io.StringIO(text),
            dialect,
        )

        rows = list(reader)

        content = "\n".join(
            " | ".join(row)
            for row in rows
        )

        return {
            "content": content,
            "rows": rows,
            "metadata": {
                "row_count": len(rows),
                "column_count": max(
                    (
                        len(row)
                        for row in rows
                    ),
                    default=0,
                ),
            },
        }

    # =========================================================
    # XLSX
    # =========================================================

    def _extract_xlsx(
        self,
        path: Path,
    ) -> dict[str, Any]:

        workbook = load_workbook(
            filename=str(path),
            read_only=True,
            data_only=True,
        )

        sheets = {}
        content_parts = []

        for worksheet in workbook.worksheets:
            rows = []

            for row in worksheet.iter_rows(
                values_only=True
            ):
                values = [
                    ""
                    if value is None
                    else str(value)
                    for value in row
                ]

                rows.append(values)

            sheets[worksheet.title] = rows

            sheet_text = "\n".join(
                " | ".join(row)
                for row in rows
            )

            content_parts.append(
                f"Sheet: {worksheet.title}\n"
                f"{sheet_text}"
            )

        content = "\n\n".join(
            content_parts
        )

        return {
            "content": content,
            "sheets": sheets,
            "metadata": {
                "sheet_count": len(
                    workbook.worksheets
                ),
                "sheet_names": [
                    worksheet.title
                    for worksheet in workbook.worksheets
                ],
                "characters": len(content),
            },
        }

    # =========================================================
    # PPTX
    # =========================================================

    def _extract_pptx(
        self,
        path: Path,
    ) -> dict[str, Any]:

        presentation = Presentation(
            str(path)
        )

        slides = []

        for slide_number, slide in enumerate(
            presentation.slides,
            start=1,
        ):
            texts = []

            for shape in slide.shapes:

                if not hasattr(
                    shape,
                    "text",
                ):
                    continue

                text = shape.text.strip()

                if text:
                    texts.append(text)

            slide_content = "\n".join(
                texts
            )

            slides.append(
                {
                    "slide": slide_number,
                    "content": slide_content,
                }
            )

        full_text = "\n\n".join(
            (
                f"Slide {slide['slide']}:\n"
                f"{slide['content']}"
            )
            for slide in slides
        )

        return {
            "content": full_text,
            "slides": slides,
            "metadata": {
                "slide_count": len(slides),
                "characters": len(full_text),
            },
        }